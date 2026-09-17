import io
import openpyxl
from app.models.models import Tournament as TournamentModel, Player as PlayerModel, PlayoffStats as PlayoffStatsModel, ConsolationStats as ConsolationStatsModel

class TournamentExportService:

    @staticmethod
    def generate_results_excel(tournament_id: int) -> io.BytesIO:
        """Vygeneruje Excel soubor s konečným pořadím turnaje."""
        tournament = TournamentModel.query.get_or_404(tournament_id)
        players = PlayerModel.query.filter_by(tournament_id=tournament_id).all()
        results_data = []

        for player in players:
            p_stats = PlayoffStatsModel.query.filter_by(player_id=player.id).first()
            c_stats = ConsolationStatsModel.query.filter_by(player_id=player.id).first()

            rank = None
            if p_stats and p_stats.final_rank is not None:
                rank = p_stats.final_rank
            elif c_stats and c_stats.final_rank is not None:
                rank = c_stats.final_rank

            if rank is not None:
                results_data.append((player, rank))

        results_data.sort(key=lambda x: x[1])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Konečné pořadí"

        # Styly
        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        align_center = openpyxl.styles.Alignment(horizontal="center", vertical="center")
        border_thin = openpyxl.styles.Border(
            left=openpyxl.styles.Side(style='thin', color='D1D5DB'),
            right=openpyxl.styles.Side(style='thin', color='D1D5DB'),
            top=openpyxl.styles.Side(style='thin', color='D1D5DB'),
            bottom=openpyxl.styles.Side(style='thin', color='D1D5DB')
        )

        # Nadpis
        ws.merge_cells("A1:B1")
        ws["A1"] = f"Výsledky turnaje: {tournament.name}"
        ws["A1"].font = openpyxl.styles.Font(size=14, bold=True)
        ws["A1"].alignment = openpyxl.styles.Alignment(horizontal="center")

        # Hlavička
        headers = ["Pořadí", "Hráč"]
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header_title
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border_thin

        # Data
        for row_idx, (player, rank) in enumerate(results_data, start=4):
            c1 = ws.cell(row=row_idx, column=1, value=f"{rank}.")
            c2 = ws.cell(row=row_idx, column=2, value=player.name)
            c1.alignment = align_center
            c1.border = border_thin
            c2.border = border_thin

        # Šířka sloupců
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_length + 5, 15)

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return file_stream